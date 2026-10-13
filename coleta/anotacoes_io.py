#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Le as anotacoes de otimizacoes/testes que o usuario registrou -- de
'anotacoes.json' e/ou 'anotacoes.xlsx' (aba 'anotacoes') na pasta do projeto --
e devolve uma lista normalizada pro painel embutir e marcar nos graficos.

Colunas/campos: site, data (AAAA-MM-DD), tipo, titulo, nota, status, data_fim.
- tipo:   otimizacao | formato_add | formato_rem | teste | outro
- status: ativo | oculto | encerrado   (nada e excluido; 'oculto' some do
          grafico mas fica no historico; 'encerrado' = teste finalizado)

Os dois formatos coexistem: o painel tambem deixa o usuario adicionar/editar
direto na tela (salvo no navegador) e exportar um 'anotacoes.json' -- que este
modulo le na proxima geracao, tornando as marcacoes permanentes.
"""
import os
import json


def _datestr(x):
    if x is None:
        return None
    if hasattr(x, "strftime"):
        return x.strftime("%Y-%m-%d")
    s = str(x).strip()
    if s.replace(".", "").isdigit() and len(s.split(".")[0]) == 8:  # 20260806
        s = s.split(".")[0]
        return f"{s[:4]}-{s[4:6]}-{s[6:]}"
    s = s[:10]
    return s if (len(s) == 10 and s[4] == "-") else None


_TIPOS = {"otimizacao", "formato_add", "formato_rem", "teste", "outro"}
_STATUS = {"ativo", "oculto", "encerrado"}


def _norm(rec, i):
    tipo = str(rec.get("tipo") or "otimizacao").strip().lower()
    status = str(rec.get("status") or "ativo").strip().lower()
    return {
        "id": str(rec.get("id") or "").strip() or ("file-" + str(i)),
        "site": str(rec.get("site") or "").strip(),
        "data": _datestr(rec.get("data")),
        "tipo": tipo if tipo in _TIPOS else "otimizacao",
        "titulo": str(rec.get("titulo") or "").strip(),
        "nota": str(rec.get("nota") or "").strip(),
        "status": status if status in _STATUS else "ativo",
        "data_fim": _datestr(rec.get("data_fim")) if rec.get("data_fim") else None,
    }


def _ler_json(path):
    if not os.path.exists(path):
        return []
    try:
        raw = json.load(open(path, encoding="utf-8"))
        if isinstance(raw, dict):
            raw = raw.get("anotacoes", [])
        return list(raw or [])
    except Exception as e:  # noqa: BLE001
        print("  ! aviso: anotacoes.json invalido, ignorando:", e)
        return []


def _ler_xlsx(path):
    if not os.path.exists(path):
        return []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        ws = wb["anotacoes"] if "anotacoes" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        hdr = [str(c).strip().lower() if c else "" for c in rows[0]]

        def gi(n):
            return hdr.index(n) if n in hdr else None

        cols = ["site", "data", "tipo", "titulo", "nota", "status", "data_fim", "id"]
        idx = {k: gi(k) for k in cols}
        out = []
        for r in rows[1:]:
            if not r:
                continue
            rec = {k: (r[idx[k]] if idx[k] is not None and idx[k] < len(r) else None) for k in cols}
            if not rec.get("data") and not rec.get("titulo"):
                continue
            out.append(rec)
        return out
    except Exception as e:  # noqa: BLE001
        print("  ! aviso: anotacoes.xlsx invalido, ignorando:", e)
        return []


def ler_anotacoes(base_dir):
    """Le anotacoes.json + anotacoes.xlsx (o que existir) e devolve lista
    normalizada (so registros com site e data validos)."""
    brutas = _ler_json(os.path.join(base_dir, "anotacoes.json")) + \
        _ler_xlsx(os.path.join(base_dir, "anotacoes.xlsx"))
    out = []
    vistos = set()
    for i, rec in enumerate(brutas):
        n = _norm(rec, i)
        if not n["site"] or not n["data"]:
            continue
        if n["id"] in vistos:
            n["id"] = n["id"] + "-" + str(i)
        vistos.add(n["id"])
        out.append(n)
    return out
