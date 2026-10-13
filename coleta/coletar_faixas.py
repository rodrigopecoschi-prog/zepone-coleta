#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
coletar_faixas.py -- lances por faixa de preco, para o simulador de regra de
preco do ZepOne Metas.

Pede ao Ad Manager, para cada rede do sites.xlsx, os lances dos ultimos 7 dias
fechados por formato, faixa de preco ($0,10) e resultado (ganhou, barrado no
piso, superado), e a regra de preco em vigor em cada formato. Grava
faixas_export.json, que o exportar_metas.py junta aos dados do painel.

So LE relatorios. Nao muda nada no Ad Manager nem no coletor.
"""
import collections
import datetime
import json
import os
import re
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:  # noqa: BLE001
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)

from gam_lib import auth, reports  # noqa: E402
from exportar_metas import grupo, NOMES  # noqa: E402

DIAS = 7          # janela, em dias fechados
TETO = 3.0        # faixas a partir deste valor ficam somadas numa so
ADX = ["AD_EXCHANGE_LINE_ITEM_LEVEL_IMPRESSIONS", "AD_EXCHANGE_LINE_ITEM_LEVEL_REVENUE"]


def _sites():
    """sites.xlsx -> {network_code: {root_ad_unit_nome minusculo: nome do site no painel}}"""
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(BASE, "sites.xlsx"), data_only=True, read_only=True)
    ws = wb["sites"] if "sites" in wb.sheetnames else wb[wb.sheetnames[0]]
    redes, hdr = collections.OrderedDict(), None
    for r in ws.iter_rows(values_only=True):
        if hdr is None:
            hdr = [str(c).strip().lower() if c else "" for c in r]
            continue
        g = lambda k: (str(r[hdr.index(k)]).strip() if k in hdr and hdr.index(k) < len(r) and r[hdr.index(k)] not in (None, "") else "")
        sid, nc = g("id"), g("network_code")
        if not sid or not nc:
            if not any(r):
                break
            continue
        raiz = (g("root_ad_unit_nome") or sid).lower()
        redes.setdefault(nc, {})[raiz] = NOMES.get(sid) or g("nome") or sid
    return redes


def _limpa(t):
    return re.sub(r"\s*\(\d+\)\s*$", "", str(t)).strip()


def _site_fmt(breadcrumb, raizes):
    partes = [_limpa(p) for p in str(breadcrumb).split("»")]
    site = raizes.get(partes[0].lower()) if partes else None
    if not site:
        return None, None
    return site, grupo(partes[-1])


def _num(v):
    try:
        return float(str(v).replace(",", "").strip() or 0)
    except Exception:  # noqa: BLE001
        return 0.0


def _col(header, nome):
    return next((k for k, h in enumerate(header) if str(h).endswith("." + nome) or str(h) == nome), None)


def main():
    with open(os.path.join(BASE, "config.json"), encoding="utf-8") as f:
        config = json.load(f)
    api = config["gam_api_version"]
    poll = int(config.get("poll_segundos", 3) or 3)
    fim = datetime.date.today() - datetime.timedelta(days=1)
    ini = fim - datetime.timedelta(days=DIAS - 1)
    # site -> formato -> faixa -> [ganhadores, receita ganhadores, barrados no piso, superados]
    faixas = collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])))
    regras = collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(float)))
    print(f"Lances por faixa de preco | {ini} a {fim}")
    for nc, raizes in _sites().items():
        try:
            client = auth.obter_client(nc, config, BASE)
            job = reports.montar_report_job(config, ini, fim, dimensions=["AD_UNIT_NAME", "BID_RANGE", "BID_REJECTION_REASON_NAME"],
                                            columns=["BID_COUNT", "BID_AVERAGE_CPM"], ad_unit_view="FLAT")
            header, linhas = reports.ler_csv(reports.rodar_report(client, api, job, poll_time_seconds=poll))
        except Exception as e:  # noqa: BLE001
            print(f"  rede {nc}: sem dados de lances ({type(e).__name__}: {str(e)[:120]})")
            continue
        ia, ib, im, ic, ip = (_col(header, n) for n in ("AD_UNIT_NAME", "BID_RANGE", "BID_REJECTION_REASON_NAME", "BID_COUNT", "BID_AVERAGE_CPM"))
        n = 0
        for ln in linhas:
            site, fmt = _site_fmt(ln[ia], raizes)
            m = re.search(r"(\d+(?:\.\d+)?)", str(ln[ib]))
            if not site or not m:
                continue
            lo = min(round(float(m.group(1)), 2), TETO)
            cnt, cpm, motivo = _num(ln[ic]), _num(ln[ip]) / 1e6, str(ln[im]).strip().lower()
            b = faixas[site][fmt][lo]
            if motivo == "winner":
                b[0] += cnt
                b[1] += cnt * cpm / 1000.0
            elif motivo == "floor":
                b[2] += cnt
            elif motivo == "outbid":
                b[3] += cnt
            n += 1
        try:
            job = reports.montar_report_job(config, ini, fim, dimensions=["AD_UNIT_NAME", "UNIFIED_PRICING_RULE_NAME"], columns=ADX, ad_unit_view="FLAT")
            header, linhas = reports.ler_csv(reports.rodar_report(client, api, job, poll_time_seconds=poll))
            ia, ir, ii = _col(header, "AD_UNIT_NAME"), _col(header, "UNIFIED_PRICING_RULE_NAME"), _col(header, ADX[0])
            for ln in linhas:
                site, fmt = _site_fmt(ln[ia], raizes)
                if site:
                    regras[site][fmt][str(ln[ir]).strip()] += _num(ln[ii])
        except Exception as e:  # noqa: BLE001
            print(f"  rede {nc}: sem o nome das regras de preco ({type(e).__name__})")
        print(f"  rede {nc}: {n} linhas de lances aproveitadas")
    saida = {"gerado": datetime.datetime.now().strftime("%d/%m/%Y %H:%M"), "de": ini.isoformat(), "ate": fim.isoformat(), "sites": {}}
    for site, fmts in faixas.items():
        for fmt, bs in fmts.items():
            linhas = [[lo, int(round(v[0])), round(v[1], 2), int(round(v[2])), int(round(v[3]))] for lo, v in sorted(bs.items()) if v[0] or v[2] or v[3]]
            if not sum(b[1] for b in linhas):
                continue
            rg = regras.get(site, {}).get(fmt, {})
            tot = sum(rg.values())
            nomes = [f"{k} ({v / tot * 100:.0f}%)" for k, v in sorted(rg.items(), key=lambda kv: -kv[1]) if tot and v / tot >= 0.05 and not k.startswith("(")][:3]
            saida["sites"].setdefault(site, {})[fmt] = {"b": linhas, **({"regra": ", ".join(nomes)} if nomes else {})}
    with open(os.path.join(BASE, "faixas_export.json"), "w", encoding="utf-8") as fj:
        json.dump(saida, fj, ensure_ascii=False)
    print(f"Arquivo: faixas_export.json ({len(saida['sites'])} sites)")


if __name__ == "__main__":
    main()
